"""
Estrae testo grezzo da un file caricato, indipendentemente dal formato.
Per le immagini (foto di bollette/multe fatte col telefono) usiamo direttamente
la vision del modello invece di un motore OCR locale: più semplice, nessuna
dipendenza di sistema, e generalmente più accurato su documenti fotografati storti.
"""

import base64
from pathlib import Path

from docx import Document
from pypdf import PdfReader
from langchain_anthropic import ChatAnthropic
from langchain_core.messages import HumanMessage

_vision_llm = ChatAnthropic(model="claude-sonnet-4-5", temperature=0)

_IMAGE_EXT = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png", ".webp": "image/webp"}


def extract_text(file_path: str) -> str:
    ext = Path(file_path).suffix.lower()

    if ext == ".pdf":
        return "\n\n".join(page.extract_text() or "" for page in PdfReader(file_path).pages)

    if ext == ".txt":
        return Path(file_path).read_text(encoding="utf-8", errors="replace")

    if ext == ".docx":
        return _extract_text_from_docx(file_path)

    if ext in _IMAGE_EXT:
        return _extract_text_from_image(file_path, _IMAGE_EXT[ext])

    raise ValueError(f"Formato non supportato: {ext}")


def _extract_text_from_docx(file_path: str) -> str:
    doc = Document(file_path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _extract_text_from_image(file_path: str, media_type: str) -> str:
    with open(file_path, "rb") as f:
        image_b64 = base64.b64encode(f.read()).decode("utf-8")

    message = HumanMessage(content=[
        {
            "type": "image",
            "source": {"type": "base64", "media_type": media_type, "data": image_b64},
        },
        {
            "type": "text",
            "text": (
                "Trascrivi fedelmente tutto il testo leggibile in questa immagine di un documento "
                "(bolletta, multa, contratto o simile). Restituisci solo il testo trascritto, "
                "senza commenti aggiuntivi."
            ),
        },
    ])
    response = _vision_llm.invoke([message])
    return response.content
